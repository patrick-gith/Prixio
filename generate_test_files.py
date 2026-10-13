"""Génère tous les fichiers de test dans ./test_files/."""
import csv
from pathlib import Path

OUT = Path("test_files")
OUT.mkdir(exist_ok=True)

# CSV simple
simple = [
    ["Designation", "Quantite", "Reference"],
    ["Clavier Logitech K380", "5", ""],
    ["Souris MX Master 3", "3", ""],
    ["9782070408504", "2", ""],
    ["bitcoin", "1", ""],
    ["HP Pavilion 15", "4", "GTIN1234567890123"],
]
with (OUT / "commande-simple.csv").open("w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows(simple)

# CSV avec point-virgule
semico = [
    ["Produit", "Qte", "Ref"],
    ["Ordinateur portable", "2", ""],
    ["Imprimante laser", "1", "9782070408504"],
    ["ethereum", "5", ""],
]
with (OUT / "commande-semicolon.csv").open("w", newline="", encoding="utf-8") as f:
    csv.writer(f, delimiter=";").writerows(semico)

# CSV avec erreurs
errors = [
    ["Designation", "Quantite", "Reference"],
    ["", "5", ""],
    ["Produit qui n'existe pas du tout", "2", ""],
    ["Clavier", "abc", ""],
    ["", "", ""],
    ["9782070408504", "0", ""],
]
with (OUT / "commande-erreurs.csv").open("w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows(errors)

# Excel
try:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Designation", "Quantite", "Reference"])
    ws.append(["Clavier Logitech K380", 5, ""])
    ws.append(["9782070408504", 2, ""])
    ws.append(["bitcoin", 1, ""])
    ws.append(["Souris MX Master 3", 3, ""])
    wb.save(OUT / "commande.xlsx")
except ImportError:
    print("openpyxl non installé, skip Excel")

# Word
try:
    from docx import Document
    doc = Document()
    doc.add_heading("Commande", 0)
    table = doc.add_table(rows=5, cols=3)
    data = [
        ["Designation", "Quantite", "Reference"],
        ["Clavier Logitech K380", "5", ""],
        ["9782070408504", "2", ""],
        ["bitcoin", "1", ""],
        ["Souris MX Master 3", "3", ""],
    ]
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            table.cell(i, j).text = val
    doc.save(OUT / "commande.docx")
except ImportError:
    print("python-docx non installé, skip Word")

print(f"✓ Fichiers générés dans {OUT}/")
for f in sorted(OUT.iterdir()):
    print(f"  - {f.name}")