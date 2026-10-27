"""Parser Word (.docx) — paragraphes + tableaux."""
from pathlib import Path

from .base import ParsedFile, ParsedRow, ParserBase


class WordParser(ParserBase):
    format = "docx"

    def parse(self, path: str | Path) -> ParsedFile:
        try:
            from docx import Document
        except ImportError:
            raise ImportError("Installe python-docx : pip install python-docx")

        doc = Document(str(path))

        # Priorité aux tableaux
        if doc.tables:
            table = max(doc.tables, key=lambda t: len(t.rows))
            return self._parse_table(table)

        # Fallback : paragraphes
        return self._parse_paragraphs(doc)

    def _parse_table(self, table) -> ParsedFile:
        rows_data = []
        for row in table.rows:
            cells = [self._clean(c.text) for c in row.cells]
            rows_data.append(cells)

        if not rows_data:
            return ParsedFile(columns=[], rows=[], format=self.format,
                              warnings=["Tableau vide."])

        header = rows_data[0]
        columns = [h or f"Col_{i+1}" for i, h in enumerate(header)]

        rows = []
        for i, raw in enumerate(rows_data[1:], start=1):
            if all(c == "" for c in raw):
                continue
            row_dict = {}
            for j, val in enumerate(raw):
                col = columns[j] if j < len(columns) else f"Col_{j+1}"
                row_dict[col] = val
            rows.append(ParsedRow(row_number=i, raw=row_dict))

        return ParsedFile(columns=columns, rows=rows, format=self.format)

    def _parse_paragraphs(self, doc) -> ParsedFile:
        rows = []
        row_n = 0
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            row_n += 1
            rows.append(ParsedRow(row_number=row_n, raw={"Texte": text}))
        return ParsedFile(
            columns=["Texte"], rows=rows, format=self.format,
            warnings=["Aucun tableau — extraction paragraphe par paragraphe."],
        )

    @staticmethod
    def _clean(value: str) -> str:
        return (value or "").replace("\n", " ").strip()